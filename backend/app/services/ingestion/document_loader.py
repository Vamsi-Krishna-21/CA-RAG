"""
Multi-format document extraction.

Supported formats:
- PDF
- DOCX
- PPT
- PPTX
- XLSX
- XLS
- CSV
- TXT
- MD
- JSON

All extractors return the common structure:

[
    {
        "page": 1,
        "text": "..."
    },
    ...
]

The existing cleaner and adaptive chunker can therefore continue
using the same interface regardless of the original file format.
"""

import csv
import json
import os
import subprocess
import tempfile

from app.services.ingestion.pdf_loader import extract_pages as extract_pdf_pages


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".ppt",
    ".pptx",
    ".xlsx",
    ".xls",
    ".csv",
    ".txt",
    ".md",
    ".json",
}


LIBREOFFICE_PATH = r"C:\Program Files\LibreOffice\program\soffice.exe"


def _extract_docx(filepath: str) -> list:
    from docx import Document

    document = Document(filepath)
    texts = []

    # Extract normal paragraphs.
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            texts.append(text)

    # Extract content from tables.
    for table in document.tables:
        for row in table.rows:
            cells = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    cells.append(cell_text)

            if cells:
                texts.append(" | ".join(cells))

    if not texts:
        return []

    return [
        {
            "page": 1,
            "text": "\n\n".join(texts),
        }
    ]


def _extract_pptx(filepath: str) -> list:
    from pptx import Presentation

    presentation = Presentation(filepath)
    pages = []

    for index, slide in enumerate(presentation.slides, start=1):
        texts = []

        for shape in slide.shapes:
            if not hasattr(shape, "text"):
                continue

            text = shape.text.strip()

            if text:
                texts.append(text)

        pages.append(
            {
                "page": index,
                "text": "\n".join(texts),
            }
        )

    return pages


def _convert_ppt_to_pptx(filepath: str) -> str:
    """
    Convert a legacy .ppt file to .pptx using LibreOffice.

    A temporary LibreOffice user profile is used for every conversion
    so that conversions launched from FastAPI do not conflict with
    another LibreOffice process.

    Returns:
        Path to the converted .pptx file.

    Raises:
        RuntimeError: If LibreOffice is unavailable or conversion fails.
    """

    if not os.path.isfile(LIBREOFFICE_PATH):
        raise RuntimeError(
            "LibreOffice was not found at: "
            f"{LIBREOFFICE_PATH}"
        )

    filepath = os.path.abspath(filepath)

    if not os.path.isfile(filepath):
        raise RuntimeError(
            f"Source PPT file does not exist: {filepath}"
        )

    with tempfile.TemporaryDirectory() as temp_directory:

        output_directory = os.path.join(
            temp_directory,
            "output",
        )

        profile_directory = os.path.join(
            temp_directory,
            "profile",
        )

        os.makedirs(output_directory, exist_ok=True)
        os.makedirs(profile_directory, exist_ok=True)

        profile_uri = (
            "file:///"
            + profile_directory.replace("\\", "/")
        )

        command = [
            LIBREOFFICE_PATH,
            "--headless",
            "--nologo",
            "--nodefault",
            "--nofirststartwizard",
            "--norestore",
            "--nolockcheck",
            f"-env:UserInstallation={profile_uri}",
            "--convert-to",
            "pptx",
            "--outdir",
            output_directory,
            filepath,
        ]

        result = subprocess.run(
            command,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            timeout=120,
        )

        converted_filename = (
            os.path.splitext(
                os.path.basename(filepath)
            )[0]
            + ".pptx"
        )

        converted_path = os.path.join(
            output_directory,
            converted_filename,
        )

        if (
            result.returncode != 0
            or not os.path.isfile(converted_path)
        ):
            error_message = (
                result.stderr.strip()
                or result.stdout.strip()
                or "Unknown LibreOffice conversion error."
            )

            raise RuntimeError(
                "Failed to convert PPT to PPTX: "
                f"{error_message}"
            )

        persistent_temp = tempfile.NamedTemporaryFile(
            suffix=".pptx",
            delete=False,
        )

        persistent_temp.close()

        with open(converted_path, "rb") as source:
            with open(
                persistent_temp.name,
                "wb",
            ) as destination:
                destination.write(source.read())

        return persistent_temp.name

def _extract_ppt(filepath: str) -> list:
    """
    Extract text from a legacy .ppt file.

    The file is converted to .pptx using LibreOffice and then
    processed by the existing PPTX extractor.
    """

    converted_path = _convert_ppt_to_pptx(filepath)

    try:
        return _extract_pptx(converted_path)
    finally:
        if os.path.exists(converted_path):
            os.remove(converted_path)


def _extract_xlsx(filepath: str) -> list:
    from openpyxl import load_workbook

    workbook = load_workbook(
        filepath,
        read_only=True,
        data_only=True,
    )

    pages = []

    for index, worksheet in enumerate(workbook.worksheets, start=1):
        rows = []

        for row in worksheet.iter_rows(values_only=True):
            values = []

            for value in row:
                if value is None:
                    values.append("")
                else:
                    values.append(str(value).strip())

            if any(values):
                rows.append(" | ".join(values))

        text = "\n".join(rows).strip()

        pages.append(
            {
                "page": index,
                "text": (
                    f"Sheet: {worksheet.title}\n\n{text}"
                    if text
                    else f"Sheet: {worksheet.title}"
                ),
            }
        )

    workbook.close()

    return pages


def _extract_xls(filepath: str) -> list:
    import xlrd

    workbook = xlrd.open_workbook(filepath)
    pages = []

    for index, worksheet in enumerate(workbook.sheets(), start=1):
        rows = []

        for row_index in range(worksheet.nrows):
            values = []

            for col_index in range(worksheet.ncols):
                value = worksheet.cell_value(row_index, col_index)

                if value is None:
                    values.append("")
                else:
                    values.append(str(value).strip())

            if any(values):
                rows.append(" | ".join(values))

        text = "\n".join(rows).strip()

        pages.append(
            {
                "page": index,
                "text": (
                    f"Sheet: {worksheet.name}\n\n{text}"
                    if text
                    else f"Sheet: {worksheet.name}"
                ),
            }
        )

    return pages


def _extract_csv(filepath: str) -> list:
    rows = []

    with open(
        filepath,
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as file:
        reader = csv.reader(file)

        for row in reader:
            values = [str(value).strip() for value in row]

            if any(values):
                rows.append(" | ".join(values))

    text = "\n".join(rows).strip()

    return [
        {
            "page": 1,
            "text": text,
        }
    ] if text else []


def _extract_text(filepath: str) -> list:
    with open(
        filepath,
        "r",
        encoding="utf-8-sig",
        errors="replace",
    ) as file:
        text = file.read().strip()

    return [
        {
            "page": 1,
            "text": text,
        }
    ] if text else []


def _extract_json(filepath: str) -> list:
    with open(
        filepath,
        "r",
        encoding="utf-8-sig",
        errors="replace",
    ) as file:
        data = json.load(file)

    text = json.dumps(
        data,
        indent=2,
        ensure_ascii=False,
    )

    return [
        {
            "page": 1,
            "text": text,
        }
    ] if text else []


def extract_document(filepath: str) -> list:
    """
    Extract text from a supported document.

    Returns:
        [
            {
                "page": int,
                "text": str
            },
            ...
        ]

    Raises:
        ValueError: If the file format is unsupported.
    """

    extension = os.path.splitext(filepath)[1].lower()

    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))

        raise ValueError(
            f"Unsupported file format: {extension or 'unknown'}. "
            f"Supported formats: {supported}"
        )

    if extension == ".pdf":
        return extract_pdf_pages(filepath)

    if extension == ".docx":
        return _extract_docx(filepath)

    if extension == ".ppt":
        return _extract_ppt(filepath)

    if extension == ".pptx":
        return _extract_pptx(filepath)

    if extension == ".xlsx":
        return _extract_xlsx(filepath)

    if extension == ".xls":
        return _extract_xls(filepath)

    if extension == ".csv":
        return _extract_csv(filepath)

    if extension in {".txt", ".md"}:
        return _extract_text(filepath)

    if extension == ".json":
        return _extract_json(filepath)

    raise ValueError(f"Unsupported file format: {extension}")